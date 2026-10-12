"""Extract text from PDF, DOCX and PPTX, then split it into token chunks."""
from io import BytesIO
from pathlib import Path
import hashlib
import tiktoken
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

ENCODING = tiktoken.get_encoding("cl100k_base")

def extract_sections(filename: str, file_bytes: bytes) -> list[dict]:
    extension = Path(filename).suffix.lower()
    sections = []

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        for page_no, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                sections.append({"text": text, "location": f"Page {page_no}"})

    elif extension == ".docx":
        doc = Document(BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        if text.strip():
            sections.append({"text": text, "location": "Word document"})

    elif extension == ".pptx":
        presentation = Presentation(BytesIO(file_bytes))
        for slide_no, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
            text = "\n".join(parts)
            if text.strip():
                sections.append({"text": text, "location": f"Slide {slide_no}"})
    else:
        raise ValueError("Unsupported file type. Upload PDF, DOCX, or PPTX.")

    if not sections:
        raise ValueError(
            "No readable text found. The file may be scanned or empty. "
            "Scanned PDFs need OCR, which is not included in this version."
        )
    return sections

def split_into_chunks(text: str, chunk_size: int = 500, overlap: int = 80) -> list[str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero.")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size.")
    tokens = ENCODING.encode(text)
    if not tokens:
        return []
    step = chunk_size - overlap
    chunks = []
    for start in range(0, len(tokens), step):
        chunk = ENCODING.decode(tokens[start:start + chunk_size]).strip()
        if chunk:
            chunks.append(chunk)
        if start + chunk_size >= len(tokens):
            break
    return chunks

def process_document(filename: str, file_bytes: bytes,
                     chunk_size: int = 500, overlap: int = 80) -> list[dict]:
    file_hash = hashlib.sha256(file_bytes).hexdigest()
    sections = extract_sections(filename, file_bytes)
    results = []
    chunk_index = 0
    for section_index, section in enumerate(sections):
        for chunk in split_into_chunks(section["text"], chunk_size, overlap):
            results.append({
                "text": chunk,
                "source_file": filename,
                "file_hash": file_hash,
                "source_location": section["location"],
                "section_index": section_index,
                "chunk_index": chunk_index,
            })
            chunk_index += 1
    if not results:
        raise ValueError("The document did not produce usable text chunks.")
    return results
